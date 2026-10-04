from pathlib import Path
import pypdfium2 as pdfium
from PIL import Image,ImageOps,ImageDraw
from docx import Document
from zipfile import ZipFile
import json
out=Path(__file__).parent
root=out.resolve().parents[1]
d=pdfium.PdfDocument(out/'template.pdf')
for start in range(0,len(d),4):
 sheet=Image.new('RGB',(1224,1584),'#aaa')
 for j in range(start,min(start+4,len(d))):
  im=d[j].render(scale=1).to_pil(); im.thumbnail((612,792))
  sheet.paste(im,((j-start)%2*612,(j-start)//2*792))
 sheet.save(out/f'template-contact-{start//4+1}.png')
print('template pages',len(d))
src=Path(r'C:\Users\drvon\AppData\Local\Temp\browser-use\exports\StudentShield Capstone-6396cf8d-1256-4289-9048-7cfb362f8ffe.docx')
doc=Document(src)
(out/'source-editable.txt').write_text('\n'.join(f'P{i} [{p.style.name}] {p.text}' for i,p in enumerate(doc.paragraphs))+'\nTABLES\n'+'\n'.join(f'T{i}\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows) for i,t in enumerate(doc.tables)),encoding='utf-8')
print('source editable',len(doc.paragraphs),len(doc.tables))
with ZipFile(root/'Capstone Documentation Template.docx') as z:
 for name in ['word/header1.xml','word/footer1.xml']:
  (out/Path(name).name).write_bytes(z.read(name))

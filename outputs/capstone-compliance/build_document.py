from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from corrections import REPLACE, OMIT
import re,json,hashlib

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
src=Document(OUT/'source-working-copy.docx')
d=Document(ROOT/'Capstone Documentation Template.docx')
body=d._element.body
for e in list(body):
 if e.tag!=qn('w:sectPr'): body.remove(e)
for name in ['Normal','No Spacing','Subheading','Heading 1','Heading 2','Heading 3','Heading 4','Caption','TOC 1','TOC 2','TOC 3','TOC 4']:
 if name not in d.styles: d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
 st.font.bold=name.startswith('Heading') or name=='Subheading'
 pf=st.paragraph_format;pf.line_spacing=1.15;pf.space_after=Pt(6);pf.space_before=Pt(0);pf.first_line_indent=Pt(0)
 if name.startswith('Heading') or name=='Subheading':
  pf.space_before=Pt(12);pf.keep_with_next=True;pf.keep_together=True
 if name.startswith('TOC'):
  st.font.bold=False;pf.space_after=Pt(0);pf.space_before=Pt(0);pf.line_spacing=1.15
  pf.left_indent=Pt(0 if name=='TOC 1' else 18 if name=='TOC 2' else 27)
normal=d.styles['Normal'];normal.paragraph_format.widow_control=True
records=[];tables=[];figures=[];used=set();section='Front matter'

def p(text='',bold=None,center=False,style=None):
 x=d.add_paragraph(style=style or 'Normal');x.add_run(text).bold=bold
 x.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
 return x

def heading(text,level=2,page=False):
 global section
 section=text
 x=p(text,style='Subheading' if level==3 and not re.match(r'^\d+\.',text) else f'Heading {level}',center=level==1)
 if page:x.paragraph_format.page_break_before=True
 records.append({'heading':text})
 return x

def blank_page(title):
 heading(title,1,True)
 x=p();x.paragraph_format.space_after=Pt(420)

def copy_p(i,style=None):
 used.add(i)
 if i in OMIT:return
 text=REPLACE.get(i,src.paragraphs[i].text).strip()
 if not text:return
 text=text.replace('facCritical','Critical').replace('text analysis of hugging faces','integration of the custom Hugging Face-hosted ML service')
 x=p(text,style=style)
 records.append({'source_paragraph':i,'section':section,'changed':i in REPLACE,'text':text})
 return x

def block(a,b,subheads=()):
 for i in range(a,b+1):
  if i in subheads:
   used.add(i);heading(REPLACE.get(i,src.paragraphs[i].text).strip(),3)
  else:copy_p(i)

def field(x,code):
 r=x.add_run();beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');r._r.append(beg)
 ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' '+code+' ';r._r.append(ins)
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
 r._r.append(OxmlElement('w:t'))
 end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)

def caption(text,kind):
 x=p(text,center=True,style='Caption');x.paragraph_format.keep_with_next=kind=='Table'
 (tables if kind=='Table' else figures).append(text)
 return x

def format_table(t,widths=None):
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or [398/len(t.columns)]*len(t.columns)
 for col,w in zip(t.columns,widths):col.width=Pt(w)
 pr=t._tbl.tblPr
 borders=pr.find(qn('w:tblBorders'))
 if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
 for n in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+n);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
 compact=len(t.rows)<=13 and sum(len(c.text) for row in t.rows for c in row.cells)<1100
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();e=OxmlElement('w:cantSplit');trpr.append(e)
  if ri==0:
   e=OxmlElement('w:tblHeader');trpr.append(e)
  for ci,c in enumerate(row.cells):
   if ci<len(widths):c.width=Pt(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for x in c.paragraphs:
    x.style=d.styles['Normal'];x.alignment=WD_ALIGN_PARAGRAPH.LEFT
    x.paragraph_format.line_spacing=1.0;x.paragraph_format.space_after=Pt(3);x.paragraph_format.space_before=Pt(3)
    x.paragraph_format.keep_with_next=compact and ri<len(t.rows)-1
    for r in x.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0);r.bold=ri==0

def table_rows(rows,title,widths=None):
 caption(f'Table {len(tables)+1}. {title}','Table')
 t=d.add_table(rows=0,cols=len(rows[0]))
 for row in rows:
  for c,text in zip(t.add_row().cells,row):c.text=str(text)
 format_table(t,widths);p()
 return t

def source_table(i,title,testing=False):
 rows=[[c.text for c in row.cells] for row in src.tables[i].rows]
 if testing:
  for row in rows[1:]:row[-1]=''
  if i==1:
   rows[-1][1]='Are incorrect passwords rejected when bcrypt verification fails?'
  if i==2:
   rows[2][1]='Does normal vent storage encrypt text with AES-256-CTR, and does the encryption-exception path expose the known Base64 fallback limitation?'
   rows[5][1]='Do context rules produce the expected reviewed category for labeled sarcasm examples?'
   rows[6][1]='Does an eligible Critical vent attempt parent SMS delivery when parent hiding is not selected?'
  if i==3:rows[7][1]='Does rejection suppress verified-incident alerts while separately applying any applicable false-report strike notification?'
  if i==5:rows[1][1]='Can the super admin create a student account with the form’s required profile fields?'
 return table_rows(rows,title,[85,244,69] if testing else None)

def image_existing(name,title,width=398):
 x=p(center=True);r=x.add_run();r.add_picture(str(ROOT/'outputs/capstone-template-revision'/name),width=Pt(width))
 x.paragraph_format.keep_with_next=True
 caption(f'Figure {len(figures)+1}. {title}','Figure')

# Title page: exact official title/names/date, fitted into the template's existing title blocks.
x=p(src.paragraphs[0].text,bold=True,center=True);x.paragraph_format.line_spacing=1.5;x.paragraph_format.space_after=Pt(84)
for i in range(5,10):x=copy_p(i);x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(0)
x.paragraph_format.space_after=Pt(56)
for i in range(13,16):x=copy_p(i);x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(0)
x.paragraph_format.space_after=Pt(56)
for i in range(19,25):x=copy_p(i);x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(0)
x.paragraph_format.space_after=Pt(64)
x=copy_p(29);x.alignment=WD_ALIGN_PARAGRAPH.CENTER
heading('APPROVAL SHEET FOR TITLE DEFENSE',1,True)
for i in [32,33,*range(35,40),41,43,44,46]:
 x=copy_p(i);x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(3)
for text in ['Mrs. Rodessa Bunag                         Mr. Joshua Pacer','Panelist 1                                         Panelist 2','Accepted and approved by:','Carlito L. Cunanan, MIT','Information Technology Department','Liezel L. Aguilar, MACDDS','School of Arts, Sciences and Technology']:
 x=p(text,center=True);x.paragraph_format.space_after=Pt(5)
used.update(range(30,59))
blank_page('ABSTRACT')
blank_page('ACKNOWLEDGMENT')
heading('TABLE OF CONTENTS',1,True);toc_list=p()
heading('LIST OF TABLES',1,True);table_list=p()
heading('LIST OF FIGURES',1,True);figure_list=p()

# Decimal numbering begins with the body; retain template header and border artwork.
front=d.sections[0]
pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'lowerRoman');pg.set(qn('w:start'),'1');front._sectPr.append(pg)
s=d.add_section(WD_SECTION_START.NEW_PAGE)
pg=s._sectPr.find(qn('w:pgNumType'))
if pg is None:pg=OxmlElement('w:pgNumType');s._sectPr.append(pg)
pg.set(qn('w:fmt'),'decimal');pg.set(qn('w:start'),'1')

heading('CHAPTER 1: INTRODUCTION',1)
heading('1.1 Project Context');block(96,98)
heading('1.2 Purpose and Description');block(100,108)
heading('1.3 Objectives of the Study');heading('1.3.1 General Objective',3);copy_p(111)
heading('1.3.2 Specific Objectives',3);block(113,118)
heading('1.4 Scope and Delimitations');heading('Scope',3);block(135,141);heading('Delimitations',3);block(143,147)
heading('1.5 Significance of the Study');block(120,132,subheads=[121,123,125,127,129,131])
heading('1.6 Definition of Terms');block(299,308)

heading('CHAPTER 2: REVIEW OF RELATED LITERATURE AND STUDIES',1,True)
heading('2.1 Related Literature (Foreign and Local)')
literature_titles=[149,153,156,160,164,167,171,176,182,186,190,194,198]
for i in range(149,202):
 if i in literature_titles:used.add(i);heading(REPLACE.get(i,src.paragraphs[i].text),3)
 else:copy_p(i)
 if i==179:heading('Understanding the Influence of Cyberbullying Among High School Students in Philippine Public Schools in the Post-Pandemic Era',3)
heading('2.2 Related Studies (Foreign and Local)')
study_titles=[204,207,211,214,217,220,224,228,231,235,239,243,247,251,255,259]
for i in range(204,263):
 if i in study_titles:used.add(i);heading(REPLACE.get(i,src.paragraphs[i].text),3)
 else:copy_p(i)
heading('2.3 Synthesis of the Reviewed Literature and Studies')
p('The reviewed literature motivates confidential reporting, contextual language analysis, differentiated screening categories, and human review. StudentShield combines these ideas in a school-focused platform for submitted vents and bullying reports. Its current design uses multilingual text screening, database-backed risk history, conditional notifications, and staff case management. These design connections do not demonstrate superior accuracy, completed user acceptance, or reduced harm; those outcomes require the planned evaluation.')

heading('CHAPTER 3: TECHNICAL BACKGROUND',1,True)
heading('3.1 Overview of Current Technologies');block(264,271)
heading('Risk Classification Engine',3);block(313,314)
table_rows([
 ['Stage','Current implementation','Interpretation'],
 ['Base score','50% distress confidence + 50% negative-emotion confidence','No 30% semantic-similarity component is used in the current base formula.'],
 ['Additive heuristics','Hopelessness +0.15; masked distress +0.14; emotional distress, isolation and self-worth +0.12 each; distortion and relationship distress +0.10; school avoidance +0.18.','These are software rules, not independently validated clinical weights.'],
 ['Compound signals','Adds 0.15 × (signal count − 1) when at least two selected indicators match.','Additive boost, not an exponential multiplier.'],
 ['Ordered overrides','Explicit self-harm, passive-disappearance ideation, finality, violence, sexual harassment, emotional and contextual rules.','Passive-disappearance matching has a 0.79 High floor before later adjustments; outputs can change with other matched rules.'],
 ['Risk history','Ordered risk labels and timestamps over 14 days; current label appended for trend and severe-recurrence evaluation.','MySQL-backed inputs replace the earlier FAISS recurrence store.'],
 ['Output bands','Low 1–30%; Medium 31–60%; High 61–85%; Critical 86–100%.','Scores are clamped to the final band and are not validated future-risk probabilities.'],
 ],'Current risk-screening stages (updated from the original metrics table)',[65,190,143])
heading('3.2 Evaluation of the Existing System (Manual or Automated)')
p('The original project purpose identifies decentralized reporting and delayed disclosure as the problem to be addressed. It motivates a centralized digital record for student concerns and staff responses. No measured baseline of SVBSM processing times, reporting rates, or intervention outcomes is provided, so a completed comparative evaluation is left unfinished.')
heading('3.3 Technical Feasibility')
heading('3.3.1 Hardware Specifications',3)
p('The following development and testing specifications are retained from the original requirements as proposed resources. They are not benchmark-validated minimums for local transformer inference or production capacity.')
block(463,475)
heading('3.3.2 Software Specifications',3);block(440,460,subheads=[441,442,449,455])
heading('3.3.3 Network/Architecture Specifications',3);copy_p(476);copy_p(327);copy_p(338)

heading('CHAPTER 4: METHODOLOGY, TESTING, AND EVALUATION',1,True)
heading('4.1 Software Development Life Cycle (SDLC Model)');copy_p(359)
image_existing('image1.png','Agile development model retained from the original paper')
block(360,383,subheads=[360,362,364,366,372,377])
heading('Implementation Plan',3);copy_p(574)
# Preserve original planned schedule, including its existing shaded cells.
caption(f'Table {len(tables)+1}. Original proposed implementation schedule (month indices)','Table')
t=d.add_table(rows=1,cols=11)
for j,c in enumerate(t.rows[0].cells):c.text='Activity / month' if j==0 else str(j)
for oldrow in src.tables[7].rows[1:]:
 cells=t.add_row().cells
 for j,(old,c) in enumerate(zip(oldrow.cells,cells)):
  c.text=old.text if j==0 else ''
  shade=old._tc.get_or_add_tcPr().find(qn('w:shd'))
  if shade is not None:c._tc.get_or_add_tcPr().append(deepcopy(shade))
format_table(t,[148]+[25]*10)
heading('4.2 Requirements Analysis');block(478,558,subheads=[478,479,485,496,507,512,524,529,531,540,541,545,548,555])
heading('4.3 System Design and Modeling')
heading('Conceptual Framework',3);block(273,297,subheads=[278,283,293,296])
heading('4.3.1 Context Diagram',3);image_existing('image5.png','Context diagram of existing vent and bullying-report flows');copy_p(341);copy_p(343)
heading('4.3.2 Use Case Diagram',3)
block(423,437,subheads=[424,429])
heading('4.3.3 Data Flow Diagram (DFD) or Activity Diagram',3);copy_p(353);copy_p(355)
heading('Existing Role Workflows',3);block(387,407,subheads=[])
heading('4.3.4 Entity Relationship Diagram (ERD)',3)
p('The earlier ERD is not reproduced as a current schema because it omits case-management and meeting tables and contains fields that differ from schema.sql. The following existing relationship explanations are retained; a corrected ERD figure remains unfinished.')
block(411,420)
heading('4.3.5 Data Dictionary',3)
p('This dictionary is derived from the checked-in schema.sql and describes the repository schema, not a live production database inspection. The controller uses virtual_meetings; an older migration separately defines video_conferences. The trajectory migration also differs from schema.sql in field width and precision, so installation state must be checked before applying schema changes.')
sql=(ROOT/'schema.sql').read_text(encoding='utf-8')
for name,defs in re.findall(r'CREATE TABLE IF NOT EXISTS (\w+) \((.*?)\n\);',sql,re.S):
 rows=[['Field','Type / constraints']]
 for line in defs.splitlines():
  line=line.strip().rstrip(',')
  if not line:continue
  if line.startswith('FOREIGN KEY'):rows.append(['Relationship',line])
  else:
   parts=line.split(None,1)
   if len(parts)==2:rows.append(parts)
 table_rows(rows,f'Data dictionary — {name}',[130,268])
heading('4.4 Software Testing Plan');copy_p(561)
p('The tables below preserve planned test scenarios. Result cells remain blank because the original paper provides no execution records. Existing repository unit tests are evidence that checks have been authored, not evidence of completed system or user-acceptance evaluation.')
heading('4.4.1 Unit Testing',3)
p('Authored checks include web-app-ci4/tests/unit/SecurityConfigTest.php and ml_service/tests/test_risk_history.py. They address selected security configuration and deterministic history behavior. No new pass/fail results are inserted into this paper.')
heading('4.4.2 Integration Testing',3);source_table(1,'Planned authentication module test scenarios',True);source_table(2,'Planned vent module test scenarios',True)
heading('4.4.3 System Testing',3)
for i,title in [(3,'bullying report'),(4,'dashboard and case management'),(5,'user management'),(6,'audit log')]:source_table(i,f'Planned {title} module test scenarios',True)
heading('4.5 User Acceptance Testing (UAT) and Evaluation');copy_p(576)
block(577,589,subheads=[577,579,581,586]);heading('Statistical Treatment of Data',3);block(592,608,subheads=[593,601]);source_table(8,'Proposed mean-score interpretation');block(610,619,subheads=[610]);source_table(9,'Proposed standard-deviation interpretation');block(621,628,subheads=[622])

heading('CHAPTER 5: RESULTS AND DISCUSSION',1,True)
heading('5.1 System Implementation');block(571,573)
heading('5.1.1 Main Interface / Dashboard',3)
# No new screenshot or implied finished screenshot documentation.
p().paragraph_format.space_after=Pt(48)
heading('5.1.2 Core Modules Features',3)
p('The current repository implements student registration with staff approval, authenticated vent and bullying submissions, automated vent screening with a PHP fallback, staff triage and case summaries, conditional SMS/email notifications, user and audit management, archive/restore operations, and one-to-one Jitsi meeting requests. These implementation observations do not constitute user-evaluation results.')
heading('5.2 Evaluation Results (ISO/IEC 25010)');p().paragraph_format.space_after=Pt(72)
heading('5.3 Discussion of Findings');p().paragraph_format.space_after=Pt(72)
heading('CONCLUSIONS AND RECOMMENDATIONS',1,True)
heading('6.1 Summary of Findings');p().paragraph_format.space_after=Pt(96)
heading('6.2 Conclusions');p().paragraph_format.space_after=Pt(96)
heading('6.3 Recommendations');p().paragraph_format.space_after=Pt(96)
heading('BACK MATTERS',1,True);heading('References (APA Format)')
refs=[REPLACE.get(i,src.paragraphs[i].text).strip() for i in range(640,669)]
refs[-1]+=' '+src.paragraphs[669].text
refs.append('Nokham, M., Plubin, B., Bunyatisai, W., Mouktonglang, T., & Plubin, S. (2025). Classifying the severity of cyberbully from social media comments. International Journal of Computer Applications, 186(63), 34–42. https://doi.org/10.5120/ijca2025924448')
for text in sorted(refs,key=str.casefold):
 x=p(text);x.paragraph_format.left_indent=Inches(.5);x.paragraph_format.first_line_indent=Inches(-.5);x.alignment=WD_ALIGN_PARAGRAPH.LEFT;x.paragraph_format.keep_together=True
 # Retain source emphasis where the reference wording is unchanged.
 for old in src.paragraphs[640:670]:
  if old.text.strip()==text and any(r.italic for r in old.runs):
   for child in list(x._p):
    if child.tag!=qn('w:pPr'):x._p.remove(child)
   cursor=0;spans=[]
   for r in old.runs:
    if not r.text:continue
    at=text.find(r.text,cursor)
    if at<0:continue
    if r.italic:spans.append((at,at+len(r.text)))
    cursor=at+len(r.text)
   cursor=0
   for start,end in spans:
    x.add_run(text[cursor:start]);x.add_run(text[start:end]).italic=True;cursor=end
   x.add_run(text[cursor:])
   break
used.update(range(640,670))
heading('Appendices',1,True)
for title in ['Appendix A: Relevant Source Code Snippets','Appendix B: User Manual','Appendix C: Survey Instruments / Questionnaires','Appendix D: Grammarian’s Certificate / Plagiarism Report']:
 blank_page(title)
heading('Supplementary Approval Form Retained from the Original Paper',2,True)
p('The original pre-oral approval form is retained below as an unsigned form. Unnamed panel positions remain blank; this form does not establish that the defense has been completed.')
heading('APPROVAL SHEET FOR PRE-ORAL DEFENSE',3)
for i in [61,62,*range(64,69),72,73,75,83,84,86,87]:
 x=copy_p(i);x.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Panelist: ____________________        Panelist: ____________________',center=True)
p('Panelist: ____________________        Panelist Chair: _______________',center=True)
blank_page('Curriculum Vitae')

# Template-order TOC includes only prescribed headings; actual page references are fields.
toc_targets=[('Title Page',d.paragraphs[0])]
for x in d.paragraphs:
 text=x.text
 if (text in ['APPROVAL SHEET FOR TITLE DEFENSE','ABSTRACT','ACKNOWLEDGMENT','TABLE OF CONTENTS','LIST OF TABLES','LIST OF FIGURES','CONCLUSIONS AND RECOMMENDATIONS','BACK MATTERS','References (APA Format)','Appendices','Curriculum Vitae'] or text.startswith('CHAPTER ') or re.match(r'^\d+\.\d+',text) or text.startswith('Appendix ')):
  toc_targets.append((text,x))
for n,(label,target) in enumerate(toc_targets):
 name=f'Section_{n}';start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(2000+n));start.set(qn('w:name'),name)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(2000+n));target._p.insert(0,start);target._p.append(end)
 x=d.add_paragraph(style='TOC 1' if label.startswith('CHAPTER ') or label in ['CONCLUSIONS AND RECOMMENDATIONS','BACK MATTERS'] else 'TOC 2')
 toc_list._p.addprevious(x._p);x.add_run(label+'\t');field(x,'PAGEREF '+name+' \\h')
 from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
 x.paragraph_format.tab_stops.add_tab_stop(Pt(398),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 x.paragraph_format.keep_with_next=False
# Accurate generated tables/figures lists: bookmark each caption and use PAGEREF.
for kind,lst,target in [('Table',tables,table_list),('Figure',figures,figure_list)]:
 for n,text in enumerate(lst,1):
  name=f'{kind}_{n}'
  for cap in d.paragraphs:
   if cap.text==text:
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(1000+(0 if kind=='Table' else 500)+n));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),start.get(qn('w:id')))
    cap._p.insert(0,start);cap._p.append(end);break
  x=d.add_paragraph();target._p.addprevious(x._p);x.add_run(text+'  ');field(x,'PAGEREF '+name+' \\h');x.paragraph_format.space_after=Pt(6)

# Keep the template's original footer line, adding page numbers in its existing empty paragraph.
foot=d.sections[0].footer
x=foot.paragraphs[1];x.alignment=WD_ALIGN_PARAGRAPH.RIGHT
x.paragraph_format.space_after=Pt(0);x.paragraph_format.space_before=Pt(0)
field(x,'PAGE')
for r in x.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
settings=d.settings.element
e=OxmlElement('w:updateFields');e.set(qn('w:val'),'true');settings.append(e)
d.core_properties.title=src.paragraphs[0].text
d.core_properties.author='Enriquez; Melgarejo; Meneses; Ramos; Urayanza'
d.core_properties.subject='Template compliance revision; supported content retained; unfinished results left blank'
d.save(ROOT/'StudentShield Capstone.docx')
(OUT/'transfer-log.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
unmapped=[{'index':i,'text':x.text} for i,x in enumerate(src.paragraphs) if x.text.strip() and i not in used]
(OUT/'unmapped-paragraphs.json').write_text(json.dumps(unmapped,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'paragraphs':len(d.paragraphs),'tables':len(d.tables),'figures':len(figures),'unmapped':len(unmapped)},indent=2))

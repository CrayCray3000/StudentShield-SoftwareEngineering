from pathlib import Path
from docx import Document
from pypdf import PdfReader
from zipfile import ZipFile
import json, importlib.util

root=Path(__file__).resolve().parents[2]
out=Path(__file__).parent
for name in ['Capstone Documentation Template.docx','Copy of CAPSTONE DRAFT PAPER.docx','StudentShield_Capstone_Documentation_Template_Final.docx']:
 d=Document(root/name)
 lines=[]
 for i,p in enumerate(d.paragraphs):
  lines.append(f'P{i} [{p.style.name}] {p.text}')
 for i,t in enumerate(d.tables):
  lines.append(f'TABLE {i}\n'+'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows))
 lines.append('SECTIONS\n'+'\n'.join(s._sectPr.xml for s in d.sections))
 (out/(name+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
 print(name,len(d.paragraphs),'paras',len(d.tables),'tables',len(d.sections),'sections')
r=PdfReader(root/'StudentShield Capstone.pdf')
(out/'original-pdf.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{p.extract_text()}' for i,p in enumerate(r.pages)),encoding='utf-8')
print('PDF',len(r.pages),'pages',r.metadata)
print('pdfium',bool(importlib.util.find_spec('pypdfium2')))
